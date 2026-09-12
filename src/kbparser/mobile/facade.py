"""Android-safe facade for the kbparser Python core.

This module is intentionally narrower than the desktop CLI. Android Phase 0
proves the embedded-Python path with formats that do not require desktop-only
tools such as LibreOffice, Tkinter, PyInstaller, or a Tesseract executable.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

_MAGIC_ZIP = b"PK\x03\x04"
_MAGIC_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_SUPPORTED_FORMATS = ("xls", "xlsx")
_OFFICE_ENGINE_PACKAGE = "com.lyomagit.kbparser.officeengine"
_OFFICE_ENGINE_ACTION = "com.lyomagit.kbparser.officeengine.CONVERT"

_DISABLED_FORMATS: dict[str, dict[str, str]] = {
    "doc": {
        "reason": "LibreOffice is provided by the Android office companion",
        "replacement": "Android office companion conversion",
    },
    "docx": {
        "reason": "DOCX mobile backend is delegated to the Android office companion",
        "replacement": "Android office companion conversion",
    },
    "pdf": {
        "reason": "desktop PDF parser requires PyMuPDF",
        "replacement": "Android office companion conversion",
    },
    "ocr": {
        "reason": "pytesseract requires an external Tesseract executable",
        "replacement": "native Android OCR bridge",
    },
}


def android_capabilities_json() -> str:
    """Return the current Android Phase 0 capability contract."""
    return _json({
        "supported_formats": list(_SUPPORTED_FORMATS),
        "disabled_formats": _DISABLED_FORMATS,
        "external_engines": {
            "office": {
                "package": _OFFICE_ENGINE_PACKAGE,
                "contract_action": _OFFICE_ENGINE_ACTION,
                "source": "CollaboraOnline/online fork",
                "role": "DOC/DOCX/PDF/RTF conversion through LibreOfficeKit",
            },
        },
        "python_entrypoint": "kbparser.mobile.facade.parse_path_json",
    })


def parse_path_json(
    path: str,
    profile: str = "balanced",
    original_source_name: str | None = None,
    original_source_format: str | None = None,
) -> str:
    """Parse an Android-cache file path and return a JSON string.

    Android code should copy the user-selected content URI into app-private
    cache first, then pass that cache path here. The result is always JSON so
    the Kotlin bridge has a stable contract for success and failure cases.
    """
    source = Path(path).expanduser()
    try:
        fmt = _detect_android_format(source)
    except Exception as exc:
        return _json({
            "status": "failed",
            "file": str(source),
            "message": f"{type(exc).__name__}: {exc}",
        })

    if fmt not in _SUPPORTED_FORMATS:
        disabled = _DISABLED_FORMATS.get(fmt, {
            "reason": f"{fmt} is not enabled in the Android Phase 0 profile",
            "replacement": "use a supported XLS/XLSX file for the feasibility spike",
        })
        return _json({
            "status": "unsupported",
            "format": fmt,
            "file": str(source),
            "message": disabled["reason"],
            "replacement": disabled["replacement"],
        })

    try:
        return _json(_parse_excel(
            source,
            fmt,
            profile,
            original_source_name=original_source_name,
            original_source_format=original_source_format,
        ))
    except Exception as exc:
        return _json({
            "status": "failed",
            "format": fmt,
            "file": str(source),
            "message": f"{type(exc).__name__}: {exc}",
        })


def parse_path_markdown(
    path: str,
    profile: str = "balanced",
    original_source_name: str | None = None,
    original_source_format: str | None = None,
) -> str:
    """Parse an Android-cache file path and return human-readable Markdown."""
    payload = json.loads(parse_path_json(
        path,
        profile=profile,
        original_source_name=original_source_name,
        original_source_format=original_source_format,
    ))
    if payload.get("status") != "success":
        return _error_markdown(payload)
    return _excel_markdown(payload)


def _detect_android_format(path: Path) -> str:
    if not path.is_file():
        raise FileNotFoundError(path)
    ext = path.suffix.lower().lstrip(".")
    with open(path, "rb") as handle:
        head = handle.read(8)
    if ext == "xlsx":
        if not head.startswith(_MAGIC_ZIP):
            raise ValueError(f"{path.name}: .xlsx but no ZIP magic")
        return ext
    if ext == "xls":
        if not head.startswith(_MAGIC_OLE):
            raise ValueError(f"{path.name}: .xls but no OLE magic")
        return ext
    return ext or "unknown"


def _parse_excel(
    path: Path,
    fmt: str,
    profile: str,
    *,
    original_source_name: str | None = None,
    original_source_format: str | None = None,
) -> dict[str, Any]:
    sheets = _parse_xlsx(path) if fmt == "xlsx" else _parse_xls(path)
    records = [
        {
            "type": "sheet_summary",
            "sheet_name": sheet["name"],
            "text": (
                f"{sheet['name']}: {sheet['non_empty_cells']} non-empty cells "
                f"across {sheet['rows']} rows x {sheet['columns']} columns"
            ),
        }
        for sheet in sheets
    ]
    source = {
        "path": str(path),
        "filename": original_source_name or path.name,
        "size_bytes": path.stat().st_size,
    }
    if original_source_name:
        source["converted_filename"] = path.name
    if original_source_format:
        source["original_format"] = original_source_format

    return {
        "status": "success",
        "format": fmt,
        "backend": "mobile-lightweight-excel",
        "profile": profile,
        "source": source,
        "sheets": sheets,
        "records": records,
    }


def _parse_xlsx(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=False, read_only=True)
    try:
        return [
            _summarize_rows(
                name=sheet.title,
                rows=sheet.iter_rows(values_only=True),
                max_row=sheet.max_row or 0,
                max_column=sheet.max_column or 0,
            )
            for sheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _parse_xls(path: Path) -> list[dict[str, Any]]:
    import xlrd

    workbook = xlrd.open_workbook(str(path), on_demand=True, use_mmap=False)
    try:
        out: list[dict[str, Any]] = []
        for sheet in workbook.sheets():
            rows = (
                tuple(sheet.cell_value(row_index, col_index) for col_index in range(sheet.ncols))
                for row_index in range(sheet.nrows)
            )
            out.append(_summarize_rows(
                name=sheet.name,
                rows=rows,
                max_row=sheet.nrows,
                max_column=sheet.ncols,
            ))
        return out
    finally:
        workbook.release_resources()


def _summarize_rows(
    *,
    name: str,
    rows,
    max_row: int,
    max_column: int,
    preview_limit: int = 10,
) -> dict[str, Any]:
    preview: list[list[Any]] = []
    non_empty = 0
    for row in rows:
        values = [_cell_json(value) for value in row]
        if any(value not in (None, "") for value in values):
            non_empty += sum(1 for value in values if value not in (None, ""))
            if len(preview) < preview_limit:
                preview.append(values)
    return {
        "name": name,
        "rows": max_row,
        "columns": max_column,
        "non_empty_cells": non_empty,
        "preview_rows": preview,
    }


def _cell_json(value):
    isoformat = getattr(value, "isoformat", None)
    if callable(isoformat):
        return isoformat()
    return value


def _json(payload: dict[str, Any]) -> str:
    return json.dumps(payload, ensure_ascii=False, sort_keys=True)


def _excel_markdown(payload: dict[str, Any]) -> str:
    source = payload.get("source", {})
    title = _clean_text(source.get("filename") or "selected-document")
    lines: list[str] = [
        f"# {title}",
        "",
        f"- Format: `{payload.get('format', 'unknown')}`",
        f"- Backend: `{payload.get('backend', 'unknown')}`",
        f"- Profile: `{payload.get('profile', 'balanced')}`",
        "",
    ]

    for sheet in payload.get("sheets", []):
        lines.extend([
            f"## Sheet: {_clean_text(sheet.get('name') or 'Sheet')}",
            "",
            (
                f"- Size: {sheet.get('rows', 0)} rows x "
                f"{sheet.get('columns', 0)} columns"
            ),
            f"- Non-empty cells: {sheet.get('non_empty_cells', 0)}",
            "",
        ])
        preview_rows = sheet.get("preview_rows") or []
        if preview_rows:
            lines.extend(_preview_table(preview_rows))
            lines.append("")

    return _normalize_blank_lines(lines).rstrip() + "\n"


def _error_markdown(payload: dict[str, Any]) -> str:
    lines = [
        "# Parse failed",
        "",
        f"- Status: `{payload.get('status', 'failed')}`",
    ]
    if payload.get("format"):
        lines.append(f"- Format: `{payload['format']}`")
    if payload.get("message"):
        lines.append(f"- Message: {_clean_text(payload['message'])}")
    if payload.get("replacement"):
        lines.append(f"- Replacement: {_clean_text(payload['replacement'])}")
    lines.append("")
    return "\n".join(lines)


def _preview_table(rows: list[list[Any]]) -> list[str]:
    width = max((len(row) for row in rows), default=0)
    if width == 0:
        return []
    header = ["Row"] + [f"Column {index + 1}" for index in range(width)]
    lines = [
        "| " + " | ".join(_escape_table_cell(value) for value in header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row_index, row in enumerate(rows, start=1):
        values = [_cell_text(value) for value in row]
        values.extend([""] * (width - len(values)))
        lines.append(
            "| "
            + " | ".join(_escape_table_cell(value) for value in [str(row_index), *values])
            + " |"
        )
    return lines


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value)


def _clean_text(value: Any) -> str:
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _escape_table_cell(value: Any) -> str:
    return _clean_text(value).replace("|", "\\|").replace("\n", "<br>")


def _normalize_blank_lines(lines: list[str]) -> str:
    normalized: list[str] = []
    previous_blank = False
    for line in lines:
        blank = line == ""
        if blank and previous_blank:
            continue
        normalized.append(line)
        previous_blank = blank
    return "\n".join(normalized)
