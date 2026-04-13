"""Generate deterministic XLSX fixture. Changing content breaks golden tests."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook


def build_basic(path: Path) -> Path:
    wb = Workbook()

    # Sheet 1: "Sales" — title block, data table, formula, notes.
    s1 = wb.active
    s1.title = "Sales"
    s1["A1"] = "Q1 Sales Report"
    s1.merge_cells("A1:D1")  # title merged row
    s1["A2"] = "Product"
    s1["B2"] = "Units"
    s1["C2"] = "Price"
    s1["D2"] = "Total"
    s1["A3"] = "Apple"
    s1["B3"] = 10
    s1["C3"] = 1.5
    s1["D3"] = "=B3*C3"
    s1["A4"] = "Pear"
    s1["B4"] = 5
    s1["C4"] = 2.0
    s1["D4"] = "=B4*C4"
    s1["A5"] = "Plum"
    s1["B5"] = 7
    s1["C5"] = 0.8
    s1["D5"] = "=B5*C5"
    s1["A7"] = "Note: preliminary figures."

    # Sheet 2: "Summary" — small table, no merges.
    s2 = wb.create_sheet("Summary")
    s2["A1"] = "Metric"
    s2["B1"] = "Value"
    s2["A2"] = "Total Units"
    s2["B2"] = 22
    s2["A3"] = "Total Revenue"
    s2["B3"] = 30.6

    # Sheet 3: hidden.
    s3 = wb.create_sheet("Archive")
    s3["A1"] = "legacy"
    s3.sheet_state = "hidden"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


if __name__ == "__main__":
    import sys
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("fixtures/xlsx/basic.xlsx")
    print(build_basic(out))
