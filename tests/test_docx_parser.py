"""Golden tests for real DOCX parser against generated fixture."""
from __future__ import annotations

from pathlib import Path

import pytest

from kbparser.parsers.base import ParseContext
from kbparser.parsers.docx import DOCXParser
from kbparser.validation import validate

from .fixtures_gen.build_docx import build_basic


@pytest.fixture(scope="module")
def basic_docx(tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("docx") / "basic.docx"
    return build_basic(out)


def _parse(path: Path):
    return DOCXParser().parse(ParseContext(path=path, profile="fidelity"))


def test_sections_tree(basic_docx: Path):
    doc = _parse(basic_docx)
    titles = [s.title for s in doc.sections]
    assert titles == ["Introduction", "Methods", "Materials", "Conclusion"]
    levels = [s.level for s in doc.sections]
    assert levels == [1, 1, 2, 1]

    by_title = {s.title: s for s in doc.sections}
    assert by_title["Introduction"].parent_id is None
    assert by_title["Materials"].parent_id == by_title["Methods"].id
    assert by_title["Conclusion"].parent_id is None


def test_blocks_and_list_items(basic_docx: Path):
    doc = _parse(basic_docx)
    types = [b.type for b in doc.blocks]
    assert types.count("heading") == 4
    assert types.count("list_item") == 3
    assert types.count("table_ref") == 2  # basic table + merged-cell table
    assert types.count("paragraph") == len(doc.blocks) - 9


def test_table_shape(basic_docx: Path):
    doc = _parse(basic_docx)
    assert len(doc.tables) == 2
    t = doc.tables[0]  # first (basic) table under Materials
    assert t.columns == ["Item", "Qty", "Price"]
    assert len(t.rows) == 3
    assert t.rows[1][0].text == "Apple"
    assert t.rows[2][0].text == "Pear"
    assert all(cell.col_span == 1 and cell.row_span == 1 for row in t.rows for cell in row)


def test_merged_table_spans(basic_docx: Path):
    doc = _parse(basic_docx)
    # The second table exercises merges: row 0 has a single cell spanning 3 cols.
    merged_t = doc.tables[1]
    assert len(merged_t.rows[0]) == 1
    title_cell = merged_t.rows[0][0]
    assert title_cell.text == "Merged Title"
    assert title_cell.col_span == 3
    # Row 2 has the vMerge anchor "A" with row_span=2; row 3 has only m2/2.
    assert merged_t.rows[2][0].text == "A"
    assert merged_t.rows[2][0].row_span == 2
    assert len(merged_t.rows[3]) == 2  # left column absorbed by the vMerge
    assert merged_t.rows[3][0].text == "m2"


def test_validation_passes(basic_docx: Path):
    doc = _parse(basic_docx)
    validate(doc)


def test_deterministic_ids(basic_docx: Path):
    a = _parse(basic_docx)
    b = _parse(basic_docx)
    assert a.id == b.id
    assert [s.id for s in a.sections] == [s.id for s in b.sections]
    assert [blk.id for blk in a.blocks] == [blk.id for blk in b.blocks]
    assert [t.id for t in a.tables] == [t.id for t in b.tables]


def test_table_linked_to_materials_section(basic_docx: Path):
    doc = _parse(basic_docx)
    materials = next(s for s in doc.sections if s.title == "Materials")
    assert doc.tables[0].section_id == materials.id


def test_root_section_fallback_when_no_heading_styles(tmp_path: Path):
    """DOCX without Word heading styles should still produce records."""
    from docx import Document as DocxDoc

    from kbparser.records import build_records

    # Build a DOCX with only 'Normal' styled paragraphs — no Heading styles.
    doc = DocxDoc()
    doc.add_paragraph("Основные положения учетной политики")
    doc.add_paragraph("I. Общие положения")
    doc.add_paragraph("Учетная политика формируется в соответствии с ПБУ 1/2008.")
    doc.add_paragraph("II. Первичные учетные документы")
    doc.add_paragraph("Бухгалтерский учет ведется автоматизированным способом.")
    p = tmp_path / "no_headings.docx"
    doc.save(str(p))

    d = _parse(p)

    # Must have at least one section (fallback root)
    assert len(d.sections) >= 1
    root = d.sections[0]
    assert root.level == 1
    # All blocks should be assigned to the root section
    assert all(b.section_id == root.id for b in d.blocks)

    # Records must be non-empty
    records = build_records(d)
    assert len(records) > 0
    assert any(r.type == "chunk" for r in records)
    # The chunk text should contain actual content
    chunks = [r for r in records if r.type == "chunk"]
    combined = " ".join(r.text for r in chunks if r.text)
    assert "учетной политики" in combined.lower() or "Общие положения" in combined


def test_preamble_and_table_before_first_heading_keep_section_and_records(tmp_path: Path):
    from docx import Document as DocxDoc

    from kbparser.records import build_records

    source = DocxDoc()
    source.add_paragraph("Important introductory conditions.")
    source.add_table(rows=1, cols=1).cell(0, 0).text = "Introductory table"
    source.add_heading("Details", 1)
    source.add_paragraph("Detailed explanation.")
    path = tmp_path / "preamble.docx"
    source.save(path)
    doc = _parse(path)
    records = build_records(doc)
    assert all(block.section_id for block in doc.blocks)
    assert doc.tables[0].section_id == doc.sections[0].id
    assert doc.sections[-1].title == "Details"
    assert any("Important introductory conditions." in (r.text or "") for r in records)
    validate(doc, records)
