from pathlib import Path

from kbparser.export import to_markdown, to_output, write_markdown
from kbparser.model import Block, Document, Parse, Source, Table, TableCell
from kbparser.versioning import PACKAGE_VERSION


def _doc() -> Document:
    source = Source(
        path="/tmp/sample.pdf",
        filename="sample.pdf",
        format="pdf",
        mime_type="application/pdf",
        sha256="0" * 64,
        size_bytes=1,
        modified_at="2026-04-14T00:00:00+00:00",
    )
    parse = Parse(
        parser="pdf",
        parser_version=PACKAGE_VERSION,
        started_at="2026-04-14T00:00:00+00:00",
        finished_at="2026-04-14T00:00:00+00:00",
    )
    return Document(
        id="doc_sample",
        source=source,
        parse=parse,
        blocks=[
            Block(id="b1", type="heading", order=0, text="Overview"),
            Block(id="b2", type="paragraph", order=1, text="Hello | world"),
            Block(id="b3", type="list_item", order=2, text="First item"),
        ],
        tables=[
            Table(
                id="t1",
                title="Scores",
                columns=["Name", "Value"],
                rows=[
                    [
                        TableCell(row=0, col=0, text="Alpha"),
                        TableCell(row=0, col=1, text="10"),
                    ],
                ],
            )
        ],
    )


def test_to_markdown_renders_readable_document():
    text = to_markdown(to_output(_doc()))

    assert text.startswith("# sample.pdf\n")
    assert "## Overview" in text
    assert "Hello \\| world" in text
    assert "- First item" in text
    assert "### Scores" in text
    assert "| Name | Value |" in text
    assert "| Alpha | 10 |" in text


def test_write_markdown_creates_utf8_file(tmp_path: Path):
    path = write_markdown(to_output(_doc()), tmp_path / "sample.pdf.md")

    assert path.exists()
    assert path.read_text(encoding="utf-8").startswith("# sample.pdf\n")


def test_docx_markdown_preserves_preamble_and_table_order(tmp_path: Path):
    from docx import Document as DocxDoc

    from kbparser.dispatcher import dispatch

    source = DocxDoc()
    source.add_paragraph("PREAMBLE")
    source.add_heading("Details", 1)
    source.add_paragraph("BEFORE_TABLE")
    table = source.add_table(rows=2, cols=2)
    for cell, text in zip(table.rows[0].cells, ["Name", "Value"], strict=True):
        cell.text = text
    table.cell(1, 0).text = "TABLE_CONTENT"
    table.cell(1, 1).text = "42"
    source.add_paragraph("AFTER_TABLE")
    path = tmp_path / "flow.docx"
    source.save(path)
    text = to_markdown(to_output(dispatch(path)))
    assert text.index("PREAMBLE") < text.index("## Details")
    assert text.index("BEFORE_TABLE") < text.index("TABLE_CONTENT") < text.index("AFTER_TABLE")
    assert text.count("| Name | Value |") == 1


def test_markdown_merged_cells_preserve_column_positions():
    doc = _doc()
    doc.tables[0] = Table(id="t1", columns=["A", "B", "C"], rows=[
        [TableCell(row=0, col=i, text=value) for i, value in enumerate(["A", "B", "C"])],
        [TableCell(row=1, col=0, text="Merged", row_span=2),
         TableCell(row=1, col=1, text="B1"), TableCell(row=1, col=2, text="C1")],
        [TableCell(row=2, col=1, text="B2"), TableCell(row=2, col=2, text="C2")],
    ])
    text = to_markdown(to_output(doc))
    assert "|  | B2 | C2 |" in text
    assert text.count("| A | B | C |") == 1


def test_workbook_tables_stay_with_their_sheets(tmp_path: Path):
    from openpyxl import Workbook

    from kbparser.dispatcher import dispatch

    book = Workbook()
    first = book.active
    first.title = "First sheet"
    second = book.create_sheet("Second sheet")
    for sheet, value in [(first, "FIRST_TABLE"), (second, "SECOND_TABLE")]:
        sheet.append(["Name", "Value"])
        sheet.append([value, 42])
    path = tmp_path / "sheets.xlsx"
    book.save(path)
    text = to_markdown(to_output(dispatch(path)))
    assert text.index("## First sheet") < text.index("FIRST_TABLE") < text.index("## Second sheet")
    assert text.index("## Second sheet") < text.index("SECOND_TABLE")


def test_failed_json_write_keeps_previous_file(tmp_path: Path, monkeypatch):
    import pytest

    from kbparser.export import write_json

    path = tmp_path / "output.json"
    path.write_text("previous valid output")
    def interrupted(payload, stream, **kwargs):
        stream.write("incomplete")
        raise OSError("disk write interrupted")
    monkeypatch.setattr("kbparser.export.json.dump", interrupted)
    with pytest.raises(OSError, match="interrupted"):
        write_json(to_output(_doc()), path)
    assert path.read_text() == "previous valid output"
    assert list(tmp_path.iterdir()) == [path]
